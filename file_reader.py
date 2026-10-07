"""
file_reader.py
Reads text from a resume file. Supports only PDF and DOCX (no OCR).
"""

import os

from docx import Document
from pypdf import PdfReader


def read_pdf(file_path):
    reader = PdfReader(file_path)
    text = ""
    for page in reader.pages:
        text += (page.extract_text() or "") + "\n"
    return text


def read_docx(file_path):
    doc = Document(file_path)
    lines = []

    # 1. normal paragraphs
    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            lines.append(paragraph.text)

    # 2. tables (education is often placed inside a table)
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                lines.append(" | ".join(cells))

    return "\n".join(lines)


def read_resume(file_path):
    """Detects .pdf or .docx and returns the extracted text."""
    extension = os.path.splitext(file_path)[1].lower()

    if extension == ".pdf":
        return read_pdf(file_path)
    elif extension == ".docx":
        return read_docx(file_path)
    else:
        raise ValueError(f"Unsupported file type: {extension} (only .pdf and .docx)")
